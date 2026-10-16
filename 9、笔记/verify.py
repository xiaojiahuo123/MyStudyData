import openpyxl, os, csv
p = r"E:\Code\MyStudyData\9、笔记\维抢修数据分析\维抢修计划表-数据分析.xlsx"
wb = openpyxl.load_workbook(p)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"  {ws.title}: {ws.max_row} 行 x {ws.max_column} 列")
ws = wb["看板"]
for r in ws.iter_rows(min_row=4, max_row=18, max_col=4, values_only=True):
    print("  KPI:", r)
csvp = r"E:\Code\MyStudyData\9、笔记\维抢修数据分析\维抢修计划表-明细.csv"
with open(csvp, encoding="utf-8-sig") as f:
    rows = list(csv.reader(f))
print("CSV 行数(含表头):", len(rows), "| 列数:", len(rows[0]))
print("CSV 表头:", rows[0])
mdp = r"E:\Code\MyStudyData\9、笔记\维抢修数据分析\周报速览.md"
md = open(mdp, encoding="utf-8").read()
print("MD 字符数:", len(md), "| 行数:", md.count(chr(10)))
print("MD 中是否出现乱码占位:", ("\ufffd" in md))
for line in md.split(chr(10))[:26]:
    print("  |", line)
